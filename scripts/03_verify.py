"""Check the Parquet against the Shapefile zips as served, by a second route, streaming.

02_export_parquet.py skips a member when its header box and then a polygon when its own box does
not touch the four prefectures. Here every polygon of every zip (the unused zips too) is tested
with `intersects` on the prepared union of the prefectures, with no box pre-filter, and the
polygons found are compared, in the same order, with the rows of the Parquet read in batches: the
source zip, member and index, the five attributes, the category and group, and the coordinates (as
WKB) must equal the source's, and no row may be left over on either side. The zips that 02 left
unused must give no polygon. Also: ids are unique, the category files hold exactly the rows of the
all-categories file for their category (ids, in order), the crs is PROJJSON, every river code
carries a river name that the source's river list gives it (its 河川名 or 河川名補足), and the
totals equal columns.json.

    uv run python scripts/03_verify.py
"""

import io
import json
import sys
from collections import Counter, defaultdict
from importlib import import_module
from pathlib import Path

import ksj_flood as k
import numpy as np
import openpyxl
import pyarrow.parquet as pq
import shapefile
import shapely
from shapely.geometry import shape

ROOT = Path(__file__).parents[1]


def river_names(xlsx: Path) -> dict[str, set[str]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    out = defaultdict(set)
    for ws in list(wb)[3:]:
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i and r[5] and r[7]:
                out[str(r[5])].add(str(r[7]))
                if r[8]:  # 河川名補足: some layers of the source use it as the river name
                    out[str(r[5])].add(str(r[8]))
    return out


def expected(run: Path, manifest: dict, area, kept_by_zip: dict):
    """Yield (zip, member, index, info, record, wkb) for every polygon that intersects `area`."""
    for name, meta in manifest["files"].items():
        if meta["role"] != "candidate":
            continue
        kept = 0
        with k.open_zip(run / name) as z:
            for member in k.shp_members(z):
                info = k.parse_member(member)
                base = member[:-4]
                sf = shapefile.Reader(
                    shp=io.BytesIO(z.read(member)), shx=io.BytesIO(z.read(base + ".shx")),
                    dbf=io.BytesIO(z.read(base + ".dbf")), encoding="cp932")  # fmt: skip
                if len(sf) == 0:
                    continue
                geoms = np.array([shape(s.__geo_interface__) for s in sf.shapes()], dtype=object)
                for i in np.nonzero(shapely.intersects(geoms, area))[0]:
                    kept += 1
                    yield name, member, int(i), info, list(sf.record(int(i))), geoms[i].wkb
        kept_by_zip[name] = kept
        print(f"{name}: {kept} polygons in the prefectures", flush=True)


def main() -> int:
    run = sorted((ROOT / "data" / "raw").glob("*/MANIFEST.json"))[-1].parent
    manifest = json.loads((run / "MANIFEST.json").read_text())
    docs = ROOT / "data" / "docs"
    par = ROOT / "data" / "parquet"
    exp = import_module("02_export_parquet")
    area = exp.prefectures([docs / Path(u).name for u in k.N03_URLS.values()])
    shapely.prepare(area)
    cols = json.loads((par / "columns.json").read_text())
    pf = pq.ParquetFile(par / "flood_area.parquet")
    crs = json.loads(pf.schema_arrow.metadata[b"geo"])["columns"]["geometry"]["crs"]
    assert isinstance(crs, dict) and crs["id"]["code"] == k.EPSG, "crs must be PROJJSON 6668"
    kept_by_zip: dict = {}
    gen = expected(run, manifest, area, kept_by_zip)
    ids: set[str] = set()
    by_cat = Counter()
    names_seen: set[tuple[str, str]] = set()
    n = 0
    for batch in pf.iter_batches(batch_size=50_000):
        d = batch.to_pydict()
        for j in range(batch.num_rows):
            try:
                zname, member, i, info, rec, wkb = next(gen)
            except StopIteration:
                raise AssertionError(
                    "the Parquet has more rows than the second route finds"
                ) from None
            assert (d["source_zip"][j], d["source_member"][j], d["source_index"][j]) == (
                zname,
                member,
                i,
            ), (n, member, i)
            got = [
                d["river_code"][j],
                d["river_name"][j],
                d["manager_code"][j],
                d["manager"][j],
                d["class_code"][j],
            ]
            assert got == [rec[0], rec[1], rec[2], rec[3], int(rec[4])], (member, i)
            assert (
                d["category"][j] == k.CATEGORIES[info["cat"]][0]
                and d["river_group"][j] == info["group"]
            )
            assert d["geometry"][j] == wkb, f"coordinates differ: {member}[{i}]"
            assert d["id"][j] not in ids, "ids are not unique"
            ids.add(d["id"][j])
            by_cat[d["category"][j]] += 1
            names_seen.add((d["river_code"][j], d["river_name"][j]))
            n += 1
    leftover = sum(1 for _ in gen)
    assert leftover == 0, f"{leftover} polygons found by the second route are not in the Parquet"
    for name in cols["unused_zips"]:
        assert kept_by_zip[name] == 0, f"{name} has polygons in the prefectures"
    assert n == cols["rows"] and dict(by_cat) == cols["by_category"], (n, by_cat)
    assert {z: v.get("polygons_kept", 0) for z, v in cols["by_zip"].items()} == kept_by_zip
    # category files: the same ids, in order, as the all-file's rows of that category
    all_ids = pq.read_table(par / "flood_area.parquet", columns=["id", "category"]).to_pydict()
    for slug in by_cat:
        t = pq.read_table(
            par / "categories" / f"{slug}.parquet", columns=["id", "category"]
        ).to_pydict()
        want = [i for i, c in zip(all_ids["id"], all_ids["category"], strict=True) if c == slug]
        assert t["id"] == want and set(t["category"]) == {slug}, slug
    names = river_names(docs / "A31_R7_datalist.xlsx")
    bad = {(c, nm) for c, nm in names_seen if nm not in names.get(c, set())}
    assert not bad, f"river names that differ from the source's river list: {sorted(bad)[:5]}"
    print(f"OK: {n} polygons; the second route finds the same set, in the same order")
    return 0


if __name__ == "__main__":
    sys.exit(main())
